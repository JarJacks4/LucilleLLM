"""
Plan documents and charts for the Claude edition.

  * PDF      — ReportLab, Escape palette, charts drawn natively on the page
  * Markdown — plain string building
  * DOCX     — python-docx (optional; a clear error if the package is absent)
  * SVG      — plain string building, returned to the app alongside the series so
               in-app charts and the PDF show the same numbers

Every chart is built from the repo: plan items actually ticked off, listening
sessions actually recorded, mood rollups, daily Self-Care Scores, active days.
When there isn't enough history the series comes back `empty` and the chart says
so — no invented numbers, ever.

Files are named `Escape-Plan-{title}-{date}.pdf` and written to GCS under
`plans/{uid}/{planId}/`, handed back as a signed URL valid for 24 h. Without a
bucket they are kept in the user's own Firestore tree and served by this API
behind the caller's token (never public).
"""

from __future__ import annotations

import base64
import html
import logging
import re
from datetime import date, datetime, timedelta
from typing import Any, Dict, List, Optional, Tuple

from escape_api import core
from escape_api.repo import Repo, user_path
from escape_api.settings import cfg

logger = logging.getLogger(__name__)

# ── Escape palette ───────────────────────────────────────────────────────────
NIGHT = "#0B1230"
LILAC = "#B9A3F0"
EMBER = "#EF7702"
MIST = "#E6EAF4"
MIST_DIM = "#8E96AE"
PALETTE = {"night": NIGHT, "lilac": LILAC, "ember": EMBER, "mist": MIST}

METRICS = ("adherence", "minutes", "mood", "score", "streak")
RANGES = {"7d": 7, "30d": 30, "90d": 90}

MIME = {"pdf": "application/pdf", "md": "text/markdown",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "svg": "image/svg+xml"}

MIN_POINTS = 2          # fewer real data points than this = honest empty state


class DocumentUnavailable(RuntimeError):
    """A format whose library is not installed on this deployment."""

    def __init__(self, fmt: str, package: str) -> None:
        super().__init__(f"{fmt.upper()} documents need the '{package}' package, which isn't installed here.")
        self.fmt = fmt
        self.package = package
        self.code = f"{fmt}_unavailable"


# ── series ───────────────────────────────────────────────────────────────────

def _days_back(tz: str, n: int) -> List[str]:
    today = date.fromisoformat(core.local_day(tz))
    return [(today - timedelta(days=i)).isoformat() for i in range(n - 1, -1, -1)]


def _label(day: str, span: int) -> str:
    d = date.fromisoformat(day)
    return d.strftime("%a") if span <= 7 else d.strftime("%d %b")


def series(repo: Repo, uid: str, tz: str, metric: str, rng: str = "30d",
           plan: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
    """
    One metric over a window, from stored data only.

    Returns {metric, range, unit, title, points:[{day,label,value|None}],
             empty, note, max, average}.
    """
    if metric not in METRICS:
        raise ValueError(f"unknown metric {metric}")
    n = RANGES.get(rng, 30)
    days = _days_back(tz, n)
    index = {d: None for d in days}
    unit, title, note = "", metric.title(), ""

    if metric == "adherence":
        unit, title = "%", "Plan adherence"
        items = (plan or {}).get("items") or []
        by_day: Dict[str, List[Dict[str, Any]]] = {}
        for i in items:
            if i.get("date") in index:
                by_day.setdefault(i["date"], []).append(i)
        today = core.local_day(tz)
        for d, day_items in by_day.items():
            if d > today:
                continue
            index[d] = round(100 * sum(1 for i in day_items if i.get("done")) / len(day_items))
        if not by_day:
            note = "This plan has no scheduled days inside that window yet."

    elif metric == "minutes":
        unit, title = "min", "Minutes of self-care"
        for _, s in repo.query(user_path(uid, "listening_sessions")):
            day = str(s.get("localDay") or (s.get("startedAt") or ""))[:10]
            if day in index and s.get("minutes"):
                index[day] = (index[day] or 0) + round(float(s["minutes"]))
        for i in ((plan or {}).get("items") or []):
            if not i.get("done") or not i.get("doneAt"):
                continue
            try:                       # doneAt is UTC; the chart is in the person's own days
                d = core.local_day(tz, core.parse_iso(str(i["doneAt"])))
            except Exception:
                d = str(i["doneAt"])[:10]
            if d in index:
                index[d] = (index[d] or 0) + int(i.get("minutes") or 0)

    elif metric == "mood":
        unit, title = "/100", "Mood"
        for _, d in repo.query(user_path(uid, "mood_days")):
            day = d.get("day")
            if day in index and d.get("count"):
                index[day] = round(100 * float(d.get("sumV", 0)) / float(d["count"]))

    elif metric == "score":
        unit, title = "/100", "Self-Care Score"
        for _, d in repo.query(user_path(uid, "scores")):
            if d.get("day") in index and d.get("score") is not None:
                index[d["day"]] = int(d["score"])

    else:  # streak
        unit, title = "days", "Streak"
        active = set(core.get_state(repo, uid).get("recentDays") or [])
        run = 0
        # walk the whole window so a streak that started before it is still right
        earliest = date.fromisoformat(days[0])
        probe = earliest - timedelta(days=1)
        while probe.isoformat() in active:
            run += 1
            probe -= timedelta(days=1)
        for d in days:
            run = run + 1 if d in active else 0
            index[d] = run
        if not active:
            index = {d: None for d in days}

    points = [{"day": d, "label": _label(d, n), "value": index[d]} for d in days]
    real = [p["value"] for p in points if p["value"] is not None]
    # A plan's scheduled days are explicit, so one of them is already worth showing.
    # Everything else needs at least two real points before a shape means anything.
    empty = len(real) < (1 if metric == "adherence" else MIN_POINTS)
    if empty and not note:
        note = ("Not enough history yet — a few more days of check-ins and this will have something "
                "honest to show.")
    return {"metric": metric, "range": rng, "unit": unit, "title": title, "points": points,
            "empty": empty, "note": note,
            # no numbers at all when the series is empty: nothing to misread
            "max": max(real) if (real and not empty) else 0,
            "average": round(sum(real) / len(real), 1) if (real and not empty) else None,
            "generatedAt": core.iso(core.now_utc())}


# ── SVG ──────────────────────────────────────────────────────────────────────

W, H = 720, 260
PAD_L, PAD_R, PAD_T, PAD_B = 46, 16, 26, 34


def _esc(s: Any) -> str:
    return html.escape(str(s), quote=True)


def chart_svg(data: Dict[str, Any]) -> str:
    """Bar chart (adherence / minutes / streak) or line (mood / score). Honest when empty."""
    head = (f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {W} {H}" width="{W}" height="{H}" '
            f'role="img" aria-label="{_esc(data["title"])}">'
            f'<rect width="{W}" height="{H}" rx="18" fill="{NIGHT}"/>'
            f'<text x="18" y="26" fill="{MIST}" font-family="Helvetica,Arial,sans-serif" '
            f'font-size="15" font-weight="600">{_esc(data["title"])}'
            f'<tspan fill="{MIST_DIM}" font-weight="400"> · last {_esc(data["range"])}</tspan></text>')
    if data.get("empty"):
        body = (f'<text x="{W // 2}" y="{H // 2}" text-anchor="middle" fill="{MIST_DIM}" '
                f'font-family="Helvetica,Arial,sans-serif" font-size="13">'
                f'{_esc(data.get("note") or "Not enough history yet.")}</text>')
        return head + body + "</svg>"

    points = data["points"]
    vals = [p["value"] for p in points if p["value"] is not None]
    top = max(vals + [1])
    if data["unit"] in ("%", "/100"):
        top = 100
    pw = W - PAD_L - PAD_R
    ph = H - PAD_T - PAD_B
    x_of = lambda i: PAD_L + (pw * (i + 0.5) / len(points))        # noqa: E731
    y_of = lambda v: PAD_T + ph - (ph * (v / top if top else 0))   # noqa: E731

    grid = "".join(
        f'<line x1="{PAD_L}" y1="{y_of(top * f):.1f}" x2="{W - PAD_R}" y2="{y_of(top * f):.1f}" '
        f'stroke="{MIST}" stroke-opacity="0.12"/>'
        f'<text x="{PAD_L - 8}" y="{y_of(top * f) + 4:.1f}" text-anchor="end" fill="{MIST_DIM}" '
        f'font-family="Helvetica,Arial,sans-serif" font-size="10">{round(top * f)}</text>'
        for f in (0, 0.5, 1.0))

    step = max(1, len(points) // 8)
    axis = "".join(
        f'<text x="{x_of(i):.1f}" y="{H - 12}" text-anchor="middle" fill="{MIST_DIM}" '
        f'font-family="Helvetica,Arial,sans-serif" font-size="10">{_esc(p["label"])}</text>'
        for i, p in enumerate(points) if i % step == 0)

    if data["metric"] in ("mood", "score"):
        seg, run = [], []
        for i, p in enumerate(points):
            if p["value"] is None:
                if len(run) > 1:
                    seg.append(run)
                run = []
            else:
                run.append(f"{x_of(i):.1f},{y_of(p['value']):.1f}")
        if len(run) > 1:
            seg.append(run)
        body = "".join(f'<polyline fill="none" stroke="{LILAC}" stroke-width="2.5" '
                       f'stroke-linejoin="round" points="{" ".join(s)}"/>' for s in seg)
        body += "".join(f'<circle cx="{x_of(i):.1f}" cy="{y_of(p["value"]):.1f}" r="3" fill="{EMBER}"/>'
                        for i, p in enumerate(points) if p["value"] is not None)
    else:
        bw = max(3.0, pw / len(points) * 0.56)
        body = ""
        for i, p in enumerate(points):
            if p["value"] is None:
                continue
            y = y_of(p["value"])
            h = max(1.5, PAD_T + ph - y)
            colour = EMBER if (data["metric"] == "adherence" and p["value"] >= 100) else LILAC
            body += (f'<rect x="{x_of(i) - bw / 2:.1f}" y="{y:.1f}" width="{bw:.1f}" height="{h:.1f}" '
                     f'rx="3" fill="{colour}"/>')

    foot = ""
    if data.get("average") is not None:
        foot = (f'<text x="{W - PAD_R}" y="26" text-anchor="end" fill="{MIST_DIM}" '
                f'font-family="Helvetica,Arial,sans-serif" font-size="11">average '
                f'{_esc(data["average"])}{_esc(data["unit"])}</text>')
    return head + grid + body + axis + foot + "</svg>"


# ── file naming ──────────────────────────────────────────────────────────────

def slug(title: str) -> str:
    s = re.sub(r"[^A-Za-z0-9]+", "-", (title or "Plan").strip()).strip("-")
    return (s[:48] or "Plan")


def document_filename(plan: Dict[str, Any], fmt: str, tz: str = "UTC") -> str:
    """Escape-Plan-{title}-{date}.pdf"""
    day = core.local_day(tz)
    return f"Escape-Plan-{slug(plan.get('title') or 'Plan')}-{day}.{fmt}"


# ── Markdown ─────────────────────────────────────────────────────────────────

DISCLAIMER = ("Made by Lucille, the AI companion in Escape. Lucille is not a therapist and this is not "
              "medical advice or a clinical plan. If you're in crisis, contact your local emergency "
              "number or a crisis line.")


def _days_grouped(plan: Dict[str, Any]) -> List[Tuple[int, str, List[Dict[str, Any]]]]:
    by_day: Dict[int, List[Dict[str, Any]]] = {}
    for i in plan.get("items", []):
        by_day.setdefault(int(i.get("day") or 1), []).append(i)
    out = []
    for day in sorted(by_day):
        items = by_day[day]
        out.append((day, items[0].get("date") or "", items))
    return out


def plan_markdown(plan: Dict[str, Any], include: Optional[List[str]] = None,
                  charts: Optional[List[Dict[str, Any]]] = None,
                  progress: Optional[Dict[str, Any]] = None) -> str:
    inc = set(include or ["schedule", "progress", "notes"])
    L: List[str] = [f"# {plan.get('title') or 'Your self-care plan'}", ""]
    if plan.get("why"):
        L += [f"_{plan['why']}_", ""]
    L += [f"- **Days:** {plan.get('days')}",
          f"- **Minutes a day:** {plan.get('minutesPerDay')}",
          f"- **Goals:** {', '.join(plan.get('goalLabels') or plan.get('goals') or []) or '—'}",
          f"- **Written by:** {'Lucille (Claude)' if plan.get('provider') == 'claude' else 'the Escape activity library'}",
          f"- **Created:** {str(plan.get('createdAt') or '')[:10]} (revision {plan.get('revision', 1)})", ""]

    if "schedule" in inc:
        L += ["## Schedule", ""]
        for day, day_date, items in _days_grouped(plan):
            L.append(f"### Day {day}" + (f" — {day_date}" if day_date else ""))
            for i in items:
                tick = "x" if i.get("done") else " "
                L.append(f"- [{tick}] **{i.get('title')}** · {i.get('minutes')} min"
                         + (f" — {i['why']}" if i.get("why") else ""))
            L.append("")

    if "progress" in inc:
        done = sum(1 for i in plan.get("items", []) if i.get("done"))
        total = len(plan.get("items", []))
        L += ["## Progress", "",
              f"- Steps done: **{done} of {total}**"
              + (f" ({round(100 * done / total)}%)" if total else ""), ""]
        if progress:
            if progress.get("score") is not None:
                L.append(f"- Self-Care Score today: **{progress['score']}** ({progress.get('band', '')})")
            if progress.get("streakDays") is not None:
                L.append(f"- Streak: **{progress['streakDays']} days**")
            L.append("")
        for ch in charts or []:
            L += [f"### {ch['title']} (last {ch['range']})", ""]
            if ch.get("empty"):
                L += [f"_{ch.get('note')}_", ""]
            else:
                L += ["| Day | " + f"{ch['title']} ({ch['unit']})".strip() + " |", "|---|---|"]
                L += [f"| {p['day']} | {'—' if p['value'] is None else p['value']} |" for p in ch["points"]]
                L.append("")

    if "notes" in inc and plan.get("notes"):
        L += ["## Notes", ""] + [f"- {str(n.get('text'))[:300]}" for n in plan["notes"]] + [""]

    L += ["---", "", DISCLAIMER, ""]
    return "\n".join(L)


# ── PDF (ReportLab) ──────────────────────────────────────────────────────────

def plan_pdf(plan: Dict[str, Any], include: Optional[List[str]] = None,
             charts: Optional[List[Dict[str, Any]]] = None,
             progress: Optional[Dict[str, Any]] = None) -> bytes:
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.platypus import Flowable, Paragraph, SimpleDocTemplate, Spacer
    except ImportError as e:  # pragma: no cover - exercised by the skip path in tests
        raise DocumentUnavailable("pdf", "reportlab") from e

    import io

    inc = set(include or ["schedule", "progress", "notes"])
    night, lilac, ember, mist = colors.HexColor(NIGHT), colors.HexColor(LILAC), colors.HexColor(EMBER), colors.HexColor(MIST)
    dim = colors.HexColor(MIST_DIM)

    def style(name, size, colour, leading=None, space=4, bold=False):
        return ParagraphStyle(name, fontName="Helvetica-Bold" if bold else "Helvetica", fontSize=size,
                              leading=leading or size * 1.35, textColor=colour, spaceAfter=space,
                              alignment=TA_LEFT)

    S = {
        "h1": style("h1", 22, mist, bold=True, space=6),
        "why": style("why", 11, lilac, space=10),
        "meta": style("meta", 9.5, dim, space=2),
        "h2": style("h2", 14, lilac, bold=True, space=6),
        "h3": style("h3", 11.5, mist, bold=True, space=3),
        "item": style("item", 10.5, mist, space=1),
        "sub": style("sub", 9, dim, space=4),
        "note": style("note", 9, dim, space=4),
    }

    class Chart(Flowable):
        """The same numbers as the SVG, drawn on the page."""

        def __init__(self, data: Dict[str, Any], width: float, height: float = 52 * mm):
            super().__init__()
            self.data, self.width, self.height = data, width, height

        def wrap(self, *_):
            return self.width, self.height

        def draw(self):
            c = self.canv
            d = self.data
            c.setFillColor(night)
            c.roundRect(0, 0, self.width, self.height, 6, stroke=0, fill=1)
            c.setFillColor(mist)
            c.setFont("Helvetica-Bold", 10)
            c.drawString(10, self.height - 16, f"{d['title']} · last {d['range']}")
            if d.get("empty"):
                c.setFillColor(dim)
                c.setFont("Helvetica", 9)
                c.drawCentredString(self.width / 2, self.height / 2 - 4,
                                    (d.get("note") or "Not enough history yet.")[:110])
                return
            pts = d["points"]
            vals = [p["value"] for p in pts if p["value"] is not None]
            top = 100 if d["unit"] in ("%", "/100") else max(vals + [1])
            x0, y0 = 34.0, 20.0
            pw, ph = self.width - x0 - 12, self.height - y0 - 26
            c.setFont("Helvetica", 7)
            for f in (0, 0.5, 1.0):
                y = y0 + ph * f
                c.setStrokeColor(colors.Color(mist.red, mist.green, mist.blue, alpha=0.18))
                c.line(x0, y, x0 + pw, y)
                c.setFillColor(dim)
                c.drawRightString(x0 - 4, y - 2, str(round(top * f)))
            line_metric = d["metric"] in ("mood", "score")
            prev = None
            bw = max(2.0, pw / len(pts) * 0.56)
            for i, p in enumerate(pts):
                x = x0 + pw * (i + 0.5) / len(pts)
                if p["value"] is None:
                    prev = None
                    continue
                y = y0 + ph * (p["value"] / top if top else 0)
                if line_metric:
                    if prev:
                        c.setStrokeColor(lilac)
                        c.setLineWidth(1.6)
                        c.line(prev[0], prev[1], x, y)
                    c.setFillColor(ember)
                    c.circle(x, y, 1.6, stroke=0, fill=1)
                    prev = (x, y)
                else:
                    c.setFillColor(ember if (d["metric"] == "adherence" and p["value"] >= 100) else lilac)
                    c.rect(x - bw / 2, y0, bw, max(0.8, y - y0), stroke=0, fill=1)
            step = max(1, len(pts) // 8)
            c.setFillColor(dim)
            c.setFont("Helvetica", 7)
            for i, p in enumerate(pts):
                if i % step:
                    continue
                c.drawCentredString(x0 + pw * (i + 0.5) / len(pts), 8, str(p["label"]))

    buf = io.BytesIO()
    title = plan.get("title") or "Your self-care plan"

    def page(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(night)
        canvas.rect(0, 0, A4[0], A4[1], stroke=0, fill=1)
        canvas.setFillColor(lilac)
        canvas.rect(0, A4[1] - 6 * mm, A4[0], 6 * mm, stroke=0, fill=1)
        canvas.setFillColor(dim)
        canvas.setFont("Helvetica", 7.5)
        canvas.drawString(18 * mm, 12 * mm, "Escape · Lucille")
        canvas.drawRightString(A4[0] - 18 * mm, 12 * mm, f"page {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(buf, pagesize=A4, title=title, author="Lucille (Escape)",
                            subject="Self-care plan", leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=20 * mm, bottomMargin=20 * mm)
    avail = doc.width
    flow: List[Any] = [Paragraph(_esc(title), S["h1"])]
    if plan.get("why"):
        flow.append(Paragraph(_esc(plan["why"]), S["why"]))
    flow += [
        Paragraph(f"{plan.get('days')} days · about {plan.get('minutesPerDay')} minutes a day", S["meta"]),
        Paragraph("Goals: " + _esc(", ".join(plan.get("goalLabels") or plan.get("goals") or []) or "—"), S["meta"]),
        Paragraph("Written by " + ("Lucille (Claude)" if plan.get("provider") == "claude"
                                   else "the Escape activity library")
                  + f" · created {str(plan.get('createdAt') or '')[:10]} · revision {plan.get('revision', 1)}",
                  S["meta"]),
        Spacer(1, 8 * mm),
    ]

    if "schedule" in inc:
        flow.append(Paragraph("Schedule", S["h2"]))
        for day, day_date, items in _days_grouped(plan):
            flow.append(Paragraph(f"Day {day}" + (f" — {day_date}" if day_date else ""), S["h3"]))
            for i in items:
                mark = "✓" if i.get("done") else "○"
                flow.append(Paragraph(f"{mark} <b>{_esc(i.get('title'))}</b> · {i.get('minutes')} min", S["item"]))
                if i.get("why"):
                    flow.append(Paragraph(_esc(i["why"]), S["sub"]))
            flow.append(Spacer(1, 3 * mm))

    if "progress" in inc:
        flow += [Spacer(1, 4 * mm), Paragraph("Progress", S["h2"])]
        done = sum(1 for i in plan.get("items", []) if i.get("done"))
        total = len(plan.get("items", []))
        flow.append(Paragraph(f"Steps done: <b>{done} of {total}</b>"
                              + (f" ({round(100 * done / total)}%)" if total else ""), S["item"]))
        if progress:
            if progress.get("score") is not None:
                flow.append(Paragraph(f"Self-Care Score today: <b>{progress['score']}</b> "
                                      f"{_esc(progress.get('band') or '')}", S["item"]))
            if progress.get("streakDays") is not None:
                flow.append(Paragraph(f"Streak: <b>{progress['streakDays']} days</b>", S["item"]))
        for ch in charts or []:
            flow += [Spacer(1, 4 * mm), Chart(ch, avail)]
            if ch.get("note") and not ch.get("empty"):
                flow.append(Paragraph(_esc(ch["note"]), S["note"]))

    if "notes" in inc and plan.get("notes"):
        flow += [Spacer(1, 4 * mm), Paragraph("Notes", S["h2"])]
        flow += [Paragraph(_esc(str(n.get("text"))[:300]), S["item"]) for n in plan["notes"]]

    flow += [Spacer(1, 8 * mm), Paragraph(_esc(DISCLAIMER), S["note"])]
    doc.build(flow, onFirstPage=page, onLaterPages=page)
    return buf.getvalue()


# ── DOCX (optional) ──────────────────────────────────────────────────────────

def plan_docx(plan: Dict[str, Any], include: Optional[List[str]] = None,
              charts: Optional[List[Dict[str, Any]]] = None,
              progress: Optional[Dict[str, Any]] = None) -> bytes:
    try:
        import io

        from docx import Document
        from docx.shared import Pt, RGBColor
    except ImportError as e:
        raise DocumentUnavailable("docx", "python-docx") from e

    inc = set(include or ["schedule", "progress", "notes"])
    d = Document()
    title = plan.get("title") or "Your self-care plan"
    h = d.add_heading(title, level=0)
    for run in h.runs:
        run.font.color.rgb = RGBColor.from_string(LILAC.lstrip("#"))
    if plan.get("why"):
        p = d.add_paragraph(plan["why"])
        p.runs[0].italic = True
    meta = d.add_paragraph()
    meta.add_run(f"{plan.get('days')} days · about {plan.get('minutesPerDay')} minutes a day\n").font.size = Pt(9)
    meta.add_run("Goals: " + (", ".join(plan.get("goalLabels") or plan.get("goals") or []) or "—") + "\n").font.size = Pt(9)
    meta.add_run("Written by " + ("Lucille (Claude)" if plan.get("provider") == "claude"
                                  else "the Escape activity library")
                 + f" · revision {plan.get('revision', 1)}").font.size = Pt(9)

    if "schedule" in inc:
        d.add_heading("Schedule", level=1)
        for day, day_date, items in _days_grouped(plan):
            d.add_heading(f"Day {day}" + (f" — {day_date}" if day_date else ""), level=2)
            for i in items:
                d.add_paragraph(f"{i.get('title')} · {i.get('minutes')} min"
                                + (f" — {i['why']}" if i.get("why") else ""),
                                style="List Bullet")

    if "progress" in inc:
        d.add_heading("Progress", level=1)
        done = sum(1 for i in plan.get("items", []) if i.get("done"))
        total = len(plan.get("items", []))
        d.add_paragraph(f"Steps done: {done} of {total}"
                        + (f" ({round(100 * done / total)}%)" if total else ""))
        if progress and progress.get("score") is not None:
            d.add_paragraph(f"Self-Care Score today: {progress['score']} {progress.get('band', '')}")
        if progress and progress.get("streakDays") is not None:
            d.add_paragraph(f"Streak: {progress['streakDays']} days")
        for ch in charts or []:
            d.add_heading(f"{ch['title']} (last {ch['range']})", level=2)
            if ch.get("empty"):
                d.add_paragraph(ch.get("note") or "Not enough history yet.")
                continue
            table = d.add_table(rows=1, cols=2)
            table.style = "Light Grid Accent 1"
            table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Day", f"{ch['title']} ({ch['unit']})".strip()
            for p in ch["points"]:
                row = table.add_row().cells
                row[0].text, row[1].text = p["day"], "—" if p["value"] is None else str(p["value"])

    if "notes" in inc and plan.get("notes"):
        d.add_heading("Notes", level=1)
        for n in plan["notes"]:
            d.add_paragraph(str(n.get("text"))[:300], style="List Bullet")

    d.add_paragraph()
    foot = d.add_paragraph(DISCLAIMER)
    foot.runs[0].font.size = Pt(8)

    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def render(fmt: str, plan: Dict[str, Any], include: Optional[List[str]] = None,
           charts: Optional[List[Dict[str, Any]]] = None,
           progress: Optional[Dict[str, Any]] = None) -> Tuple[bytes, str]:
    """(bytes, mime) for 'pdf' | 'md' | 'docx'."""
    if fmt == "pdf":
        return plan_pdf(plan, include, charts, progress), MIME["pdf"]
    if fmt == "md":
        return plan_markdown(plan, include, charts, progress).encode("utf-8"), MIME["md"]
    if fmt == "docx":
        return plan_docx(plan, include, charts, progress), MIME["docx"]
    raise ValueError(f"unsupported format {fmt}")


# ── storage + signed URLs ────────────────────────────────────────────────────

DOCS = "plan_docs"
INLINE_MAX = 6 * 1024 * 1024        # don't park anything bigger than this in Firestore


def _gcs_signed(uid: str, plan_id: str, name: str, data: bytes, mime: str, ttl_h: int) -> Optional[str]:
    bucket_name = cfg().plan_docs_bucket
    if not bucket_name:
        return None
    try:
        from google.cloud import storage as gcs
        blob = gcs.Client().bucket(bucket_name).blob(f"plans/{uid}/{plan_id}/{name}")
        blob.upload_from_string(data, content_type=mime)
        return blob.generate_signed_url(expiration=timedelta(hours=ttl_h), method="GET", version="v4")
    except Exception as e:
        logger.warning("plan document upload to gs://%s failed: %s", bucket_name, e)
        return None


def store_document(repo: Repo, uid: str, plan: Dict[str, Any], fmt: str, data: bytes, mime: str,
                   tz: str = "UTC") -> Dict[str, Any]:
    """
    Write the document and return {id, url, mime, bytes, expiresAt, filename, storage}.
    GCS + signed URL when a bucket is configured; otherwise kept in the user's own
    tree and served by GET /v2/plans/{planId}/documents/{docId} behind their token.
    """
    name = document_filename(plan, fmt, tz)
    ttl = max(1, cfg().plan_docs_ttl_hours)
    doc_id = core.new_id("doc")
    expires = core.iso(core.now_utc() + timedelta(hours=ttl))
    url = _gcs_signed(uid, plan["id"], name, data, mime, ttl)
    record = {"id": doc_id, "planId": plan["id"], "filename": name, "mime": mime,
              "bytes": len(data), "createdAt": core.iso(core.now_utc()), "expiresAt": expires,
              "format": fmt, "storage": "gcs" if url else "api"}
    if url:
        record["url"] = url
    else:
        if len(data) > INLINE_MAX:
            raise ValueError("Document is too large to serve without a GCS bucket configured.")
        record["data"] = base64.b64encode(data).decode("ascii")
        record["url"] = f"{cfg().public_base_url.rstrip('/')}/v2/plans/{plan['id']}/documents/{doc_id}"
    repo.set(user_path(uid, DOCS, doc_id), record)
    return {k: v for k, v in record.items() if k != "data"}


def read_document(repo: Repo, uid: str, doc_id: str) -> Optional[Tuple[bytes, str, str]]:
    """(data, mime, filename) for an api-served document, or None."""
    rec = repo.get(user_path(uid, DOCS, doc_id))
    if not rec or not rec.get("data"):
        return None
    try:
        when = core.parse_iso(rec["expiresAt"])
    except Exception:
        when = None
    if when and core.now_utc() > when:
        return None
    return base64.b64decode(rec["data"]), rec.get("mime") or "application/octet-stream", rec.get("filename") or doc_id


def expired(rec: Dict[str, Any], now: Optional[datetime] = None) -> bool:
    try:
        return (now or core.now_utc()) > core.parse_iso(rec["expiresAt"])
    except Exception:
        return False
